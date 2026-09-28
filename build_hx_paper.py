from __future__ import annotations

import csv
import math
from collections import Counter, defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "artifacts"
ASSET_DIR = OUT_DIR / "hx_paper_assets"
OUT_DOCX = OUT_DIR / "HX_Hybrid_Directional_Specification_Test.docx"
RANDOM_CSV = HERE / "random_1000.csv"

NAVY = "17365D"
BLUE = "2F5597"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
PALE_GRAY = "F6F6F6"
WHITE = "FFFFFF"
BLACK = "000000"
GREEN = "548235"
RED = "C00000"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7B7B7", size=4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, rows, widths=None, font_size=8.6, align=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, text in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        run.bold = True
        run.font.color.rgb = rgb(WHITE)
        run.font.size = Pt(font_size)
    for i, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            if i % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = align[j] if align else (WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
        if widths:
            for j, width in enumerate(widths):
                row.cells[j].width = width
    if widths:
        for j, width in enumerate(widths):
            hdr.cells[j].width = width
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run(text)


def add_equation(doc, text: str, number: int | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    math_para = OxmlElement("m:oMathPara")
    math_obj = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text if number is None else f"{text}                                      ({number})"
    math_run.append(math_text)
    math_obj.append(math_run)
    math_para.append(math_obj)
    p._p.append(math_para)


def add_field(run, instruction: str) -> None:
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_bullet(doc, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)


def add_number(doc, text: str, level: int = 0) -> None:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)


def add_para(doc, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = rgb(BLACK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    # The default Word template gives Title a blue bottom border. Academic
    # title pages in this manuscript use typography only, so remove it.
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles.add_style("Paper Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = rgb(GRAY)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)

    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        st = styles[name]
        st.font.name = "Aptos Display"
        st.font.color.rgb = rgb(BLACK)
        st.font.bold = True
        st.font.size = Pt(size)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(5)

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = rgb(GRAY)
    caption.paragraph_format.space_after = Pt(4)

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        st = styles[name]
        st.font.name = "Aptos"
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.15


def setup_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)


def add_header_footer(section) -> None:
    def fill_footer(footer) -> None:
        p = footer.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run()
        r.font.size = Pt(8)
        add_field(r, "PAGE")

    # Keep the manuscript header empty. This is stable across Word versions
    # and leaves the compact page number as the only running furniture.
    section.header.paragraphs[0].clear()
    fill_footer(section.footer)


def font(size=28, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\aptos.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    if bold:
        candidates = [Path(r"C:\Windows\Fonts\aptos-bold.ttf"), Path(r"C:\Windows\Fonts\arialbd.ttf")] + candidates
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def load_random_data():
    rows = []
    with RANDOM_CSV.open("r", encoding="utf-8-sig", newline="") as f:
        rows.extend(csv.DictReader(f))
    return rows


def metric_summary(rows):
    methods = {"form_check": "p_fc", "MD4": "p_md4", "HX": "p_hx", "RESET": "p_reset"}
    out = {}
    for label, col in methods.items():
        h0 = [r for r in rows if r["klass"] == "H0" and not r.get("fail_" + col[2:], "")]
        h1 = [r for r in rows if r["klass"] == "H1" and not r.get("fail_" + col[2:], "")]
        fp = sum(float(r[col]) <= 0.05 for r in h0)
        tp = sum(float(r[col]) <= 0.05 for r in h1)
        out[label] = {"fp": fp, "n0": len(h0), "tp": tp, "n1": len(h1), "fpr": fp / len(h0), "tpr": tp / len(h1)}
    return out


def wilson(k, n, z=1.959963984540054):
    p = k / n
    den = 1 + z * z / n
    mid = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return mid - half, mid + half


def draw_performance_chart(summary, path: Path):
    W, H = 1500, 770
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    f_title, f_axis, f_small = font(38, True), font(26), font(22)
    d.text((60, 26), "Randomized stress test performance", fill="#17365D", font=f_title)
    names = list(summary)
    colors = ["#A5A5A5", "#5B9BD5", "#17365D", "#ED7D31"]
    panels = [(70, 120, 700, 660, "False-positive rate", "fpr", 0.10), (800, 120, 1430, 660, "True-positive rate", "tpr", 1.00)]
    for x0, y0, x1, y1, label, key, ymax in panels:
        d.rectangle((x0, y0, x1, y1), outline="#B7B7B7", width=2)
        d.text((x0, y0 - 42), label, fill="#000000", font=f_axis)
        for tick in range(6):
            value = ymax * tick / 5
            y = y1 - (y1 - y0) * tick / 5
            d.line((x0, y, x1, y), fill="#E7E6E6", width=2)
            d.text((x0 - 58, y - 12), f"{value:.2f}", fill="#666666", font=f_small)
        if key == "fpr":
            y = y1 - (y1 - y0) * 0.05 / ymax
            d.line((x0, y, x1, y), fill="#C00000", width=4)
            d.text((x1 - 150, y - 31), "α = 0.05", fill="#C00000", font=f_small)
        gap = (x1 - x0) / len(names)
        for i, (name, color) in enumerate(zip(names, colors)):
            v = summary[name][key]
            bw = gap * 0.50
            bx0 = x0 + gap * i + gap * 0.25
            bx1 = bx0 + bw
            by = y1 - (y1 - y0) * v / ymax
            d.rectangle((bx0, by, bx1, y1), fill=color)
            d.text((bx0 + 2, by - 32), f"{v:.3f}", fill="#000000", font=f_small)
            tw = d.textbbox((0, 0), name, font=f_small)[2]
            d.text(((bx0 + bx1 - tw) / 2, y1 + 12), name, fill="#000000", font=f_small)
    im.save(path, dpi=(180, 180))


def draw_workflow(path: Path):
    W, H = 1500, 650
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    f_title, f_box, f_small = font(34, True), font(24, True), font(21)
    d.text((50, 25), "HX analysis pipeline", fill="#17365D", font=f_title)
    boxes = [
        (50, 150, 245, 330, "1  Fit base", "log-linear OLS\nand residuals"),
        (300, 150, 495, 330, "2  Generate", "structural +\nfitted-value\ndirections"),
        (550, 150, 745, 330, "3  Residualize", "against base\nmodel space"),
        (800, 150, 995, 330, "4  Select", "rank-aware\northogonal pool"),
        (1050, 105, 1250, 285, "5a  maxT", "strongest\ndirection"),
        (1050, 350, 1250, 530, "5b  Joint F", "distributed\nmisspecification"),
        (1305, 225, 1470, 405, "6  Combine", "Bonferroni\nminimum p"),
    ]
    for i, (x0, y0, x1, y1, heading, body) in enumerate(boxes):
        fill = "#EEF5FB" if i < 4 else ("#D9EAF7" if i < 6 else "#17365D")
        d.rounded_rectangle((x0, y0, x1, y1), radius=16, fill=fill, outline="#5B9BD5", width=3)
        col = "white" if i == 6 else "#17365D"
        d.text((x0 + 16, y0 + 18), heading, fill=col, font=f_box)
        for k, line in enumerate(body.split("\n")):
            d.text((x0 + 16, y0 + 68 + 34 * k), line, fill=("white" if i == 6 else "#333333"), font=f_small)
    arrows = [((245, 240), (300, 240)), ((495, 240), (550, 240)), ((745, 240), (800, 240)), ((995, 240), (1050, 195)), ((995, 240), (1050, 440)), ((1250, 195), (1305, 285)), ((1250, 440), (1305, 345))]
    for (x0, y0), (x1, y1) in arrows:
        d.line((x0, y0, x1, y1), fill="#5B9BD5", width=5)
        ang = math.atan2(y1 - y0, x1 - x0)
        for off in (2.6, -2.6):
            d.line((x1, y1, x1 + 16 * math.cos(ang + off), y1 + 16 * math.sin(ang + off)), fill="#5B9BD5", width=5)
    d.text((50, 575), "Decision rule: reject H0 when pHX ≤ α; report the winning direction only as an exploratory diagnostic.", fill="#333333", font=f_small)
    im.save(path, dpi=(180, 180))


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    rows = load_random_data()
    summary = metric_summary(rows)
    perf_path = ASSET_DIR / "randomized_performance.png"
    flow_path = ASSET_DIR / "hx_workflow.png"
    draw_performance_chart(summary, perf_path)
    draw_workflow(flow_path)

    doc = Document()
    setup_styles(doc)
    # Use the default running header on every page.
    doc.settings.odd_and_even_pages_header_footer = False
    for section in doc.sections:
        setup_page(section)
        add_header_footer(section)

    props = doc.core_properties
    props.title = "HX A Hybrid Directional Specification Test for Log Linear Physical Model Discovery"
    props.subject = "Method development and benchmark evidence"
    props.keywords = "specification testing, log-linear models, permutation inference, RESET, symbolic regression, falsification"

    # Title page
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(90)
    p.add_run("HX A Hybrid Directional Specification Test for Log Linear Physical Model Discovery")
    p = doc.add_paragraph(style="Paper Subtitle")
    p.add_run("Method Development and Benchmark Evidence")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("Anonymous manuscript")
    r.font.size = Pt(11)
    r.font.color.rgb = rgb(GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("5 September 2026")
    r.font.size = Pt(10)
    r.font.color.rgb = rgb(GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Research article")
    r.bold = True
    r.font.color.rgb = rgb(NAVY)
    doc.add_page_break()

    doc.add_heading("Abstract", level=1)
    add_para(doc, "Scientific model discovery requires more than predictive fit: a candidate law can interpolate a sample while silently violating its claimed functional form. We propose HX, a hybrid directional specification test for log-linear physical models. HX augments a restricted log-linear regression with two complementary direction families. Structural directions—quadratic, cubic, interaction, and oscillatory terms in standardized log predictors—target interpretable departures. Fitted-value directions—powers of the restricted fitted response—provide an omnibus response-sensitive probe related to Ramsey’s regression specification error test. A rank-aware modified Gram–Schmidt step removes directions already represented by the base model and limits the effective augmentation. Two residual-resampling statistics are then evaluated: a familywise maxT statistic for a concentrated departure and a joint partial-F statistic for a distributed departure. Their p-values are combined by a two-test Bonferroni minimum, so no independence assumption is required for the combination. The empirical record comprises a controlled development benchmark and a post-development randomized stress test. In 1,000 randomized groups (393 null and 607 alternative cases), HX attained a false-positive rate of 0.0509 (95% Wilson interval 0.0332–0.0773) and power of 0.9522 (0.9322–0.9665), compared with power 0.4942 for a single-direction form check, 0.8913 for the structural-only MD4 library, and 0.7512 for RESET. The randomized suite is a stress test rather than independent external validation, and frozen fitted-value directions make the reported residual-permutation calibration approximate. The method is therefore positioned as a falsification screen with explicit assumptions, failure states, and a confirmatory resampling protocol, not as a proof that a physical law is correct.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.add_run("Keywords: ").bold = True
    p.add_run("model specification; log-linear physical laws; permutation inference; RESET; symbolic regression; falsification; dimensionless modeling")

    doc.add_heading("Plain Language Summary", level=1)
    add_para(doc, "A compact scientific equation may fit data well for the wrong reason. HX asks whether the residuals contain a recognizable shape that the proposed equation should already have explained. It searches several physically interpretable shapes and two broad fitted-value shapes, protects the directional search with a maxT correction, and also checks whether the directions matter jointly. A small HX p-value means the proposed form should be challenged; a large p-value means only that this particular battery did not expose a defect.")

    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    add_table(doc, ["Section", "Purpose"], [
        ("1–3", "Motivation, related work, and problem formulation"),
        ("4–5", "HX algorithm, inferential interpretation, and safeguards"),
        ("6–7", "Evaluation protocol and empirical results"),
        ("8–11", "Mechanisms, limitations, reproducibility, and conclusion"),
        ("Appendices", "Pseudocode, failure taxonomy, and reporting checklist"),
    ], widths=[Inches(1.25), Inches(5.8)], font_size=9.2)
    doc.add_page_break()

    doc.add_heading("1 Introduction", level=1)
    add_para(doc, "Log-linear regression is a natural representation for monomial laws, scaling relations, and dimensionally motivated models. If positive variables satisfy y = c∏xⱼ^βⱼ, taking logarithms yields an additive linear equation. This convenience can obscure a central scientific question: does the assumed functional form deserve to survive, or does the residual structure contradict it? Predictive accuracy alone cannot answer that question because flexible noise, limited dynamic range, and parameter compensation can make a wrong form look adequate in sample.")
    add_para(doc, "Specification tests address a narrower and more defensible objective. They do not certify a law; they attempt to falsify a proposed form by asking whether omitted structure remains. Classical RESET uses powers of fitted values as generic probes. Domain-oriented form checks instead use targeted transformations, which are interpretable but can be blind to alternatives outside the chosen direction. The motivating benchmark behind this work showed that neither strategy uniformly dominates: structural checks can expose departures that an omnibus fitted-value test misses, while fitted-value terms can recover alternatives that a restricted structural library does not span.")
    add_para(doc, "HX turns that complementarity into one auditable procedure. It combines structural and fitted-value directions, removes redundant columns, tests both the strongest coordinate and the augmentation as a whole, and returns explicit numerical failure states rather than silently coercing ill-posed fits into a p-value. The design target is a practical falsification layer that can sit between symbolic model generation and scientific interpretation.")
    add_para(doc, "This article makes four contributions. First, it gives a complete mathematical specification of the HX statistic and its output contract. Second, it separates search multiplicity, rank deficiency, and resampling assumptions that are often conflated in informal implementations. Third, it reports controlled and randomized evidence with exact denominators, uncertainty intervals, family-level results, and paired discordance counts. Fourth, it states what the evidence does not establish and defines a confirmatory protocol for independent evaluation.")

    doc.add_heading("2 Related Work and Test Taxonomy", level=1)
    add_para(doc, "Ramsey’s RESET test augments a restricted regression with nonlinear functions of its fitted values and evaluates whether the augmentation contributes explanatory power [1]. RESET is attractive because it requires no explicit alternative, but the fitted-value basis can compress distinct physical departures into similar directions and may have low power against cancellation patterns. Heteroskedasticity-consistent covariance estimators improve robustness for Wald-style implementations but do not transform a generic basis into an interpretable mechanism [2].")
    add_para(doc, "Permutation inference provides finite-sample calibration when the chosen exchangeability conditions are satisfied. Freedman and Lane proposed a residual-based construction for general linear models [3], and Winkler et al. systematized permutation strategies, exchangeability blocks, nuisance adjustment, and error control [4]. For heteroskedastic regression errors, wild-bootstrap methods provide a more credible alternative to unrestricted residual permutation [5]. Westfall and Young’s maxT framework motivates simultaneous adjustment across correlated directional statistics [6].")
    add_para(doc, "A separate literature seeks physical structure during model discovery. Dimensionally consistent learning and Buckingham-Pi transformations reduce the search space and improve interpretability [7]. Symbolic-regression benchmarks such as SRBench, AI Feynman, PySR, and SRSD emphasize equation recovery, predictive performance, dimensional consistency, and dataset design [8–11]. HX addresses a complementary stage: after a candidate equation has been proposed and transformed to a restricted log-linear representation, does an explicit battery expose remaining structure?")
    add_para(doc, "The distinction is important. Equation discovery proposes; specification testing challenges. A discovery algorithm may rank HX-passing models above equally predictive HX-failing models, but HX does not identify the true law and should not be optimized on the same final validation data used to claim scientific adequacy.")

    doc.add_heading("3 Problem Formulation", level=1)
    add_para(doc, "Let yᵢ > 0 be a response and xᵢⱼ > 0 be p predictors for observations i = 1,…,n. Define ℓyᵢ = log yᵢ and Xℓ,ᵢⱼ = log xᵢⱼ. The restricted model contains an intercept and the standardized log predictors Z. With B = [1, Z], the null model is")
    add_equation(doc, "H₀: E(ℓy | Xℓ) ∈ span(B),     ℓy = Bβ + ε.", 1)
    add_para(doc, "The alternative is deliberately composite: at least one admissible nonlinear direction explains residual variation beyond span(B). The test is conditional on the supplied data matrix, positivity and finite-value checks, the selected direction library, and the resampling scheme. Rejection therefore means that the specified log-linear form is incompatible with at least one tested departure at the chosen level. Non-rejection means only that the battery found insufficient evidence.")
    add_table(doc, ["Symbol", "Definition"], [
        ("B", "Restricted design matrix [1, Z]"),
        ("ŷ and r", "Restricted fitted response and residual vector"),
        ("D", "Candidate augmentation matrix before residualization"),
        ("W", "Selected orthonormal augmentation coordinates"),
        ("q", "Number of retained independent directions"),
        ("Tmax", "Maximum absolute studentized directional score"),
        ("Fjoint", "Partial-F statistic for the full selected augmentation"),
        ("pHX", "Bonferroni-combined HX p-value"),
    ], widths=[Inches(1.1), Inches(5.95)], font_size=9.2)

    doc.add_heading("4 HX Core Algorithm", level=1)
    add_para(doc, "HX is designed around complementary failure modes. A concentrated defect may align strongly with one interpretable direction, whereas a diffuse defect may be weak in every coordinate but large jointly. HX therefore retains two statistics rather than forcing one representation to serve both roles.")
    doc.add_picture(str(flow_path), width=Inches(7.0))
    add_caption(doc, "Figure 1. HX analysis pipeline. The maxT and joint-F branches are calibrated under the same declared resampling policy and combined only after each p-value is obtained.")

    doc.add_heading("4.1 Restricted fit and preprocessing", level=2)
    add_para(doc, "Each log predictor is centered and scaled using the analysis sample. Constant or numerically degenerate columns are rejected before fitting. The restricted ordinary least-squares fit yields β̂ = B⁺ℓy, ŷ = Bβ̂, and r = ℓy − ŷ, where B⁺ denotes the Moore–Penrose inverse. The residual degrees of freedom must exceed a declared floor. A degenerate-fit guard is expressed as the translation-invariant ratio RSSr/TSSℓy, not as RSSr divided by the raw squared norm of ℓy.")
    add_equation(doc, "RSSᵣ = rᵀr,     TSSℓy = (ℓy − mean(ℓy)1)ᵀ(ℓy − mean(ℓy)1).", 2)

    doc.add_heading("4.2 Candidate direction library", level=2)
    add_para(doc, "For each standardized predictor zⱼ, the structural library may contain centered quadratic zⱼ², centered cubic zⱼ³, and an oscillatory probe sin(2πzⱼ). Pairwise interactions zⱼzₖ are included for j < k. The fitted-value library contains centered powers of a standardized restricted fitted vector, typically ŷs² and ŷs³, where ŷs = (ŷ − mean ŷ)/sd(ŷ). Standardizing before taking powers makes the library insensitive to an additive shift or nonzero scalar rescaling of ℓy, apart from a harmless sign change in odd powers.")
    add_equation(doc, "D = [ {zⱼ², zⱼ³, sin(2πzⱼ)}ⱼ ; {zⱼzₖ}ⱼ<ₖ ; ŷs² ; ŷs³ ].", 3)
    add_para(doc, "The names attached to retained columns are part of the diagnostic record. They support mechanistic follow-up but do not convert the selected coordinate into a separately confirmed discovery.")

    doc.add_heading("4.3 Base residualization and rank-aware selection", level=2)
    add_para(doc, "Every candidate direction is residualized against the restricted model: D⊥ = (I − PB)D, with PB = BB⁺. A modified Gram–Schmidt process selects a numerically independent set. At each step, the candidate with the largest residual norm after projection on the already selected W is retained if the norm exceeds a tolerance. If a hard direction budget is active, selection should use a relative residual norm, ‖dⱼ,res‖/max(‖dⱼ,⊥‖,τ), to avoid a scale-dependent preference. In the reported implementation no hard budget bound the evaluated low-dimensional cases; the relative criterion is the normative rule for larger libraries.")
    add_equation(doc, "wₖ = dⱼ,res / ‖dⱼ,res‖₂,     WᵀW ≈ I,     WᵀB ≈ 0.", 4)

    doc.add_heading("4.4 Directional maxT branch", level=2)
    add_para(doc, "For a retained orthonormal direction wⱼ, the score wⱼᵀr measures alignment between the restricted residual and that departure. It is studentized using the restricted residual variance estimate. HX records the maximum absolute score across q directions.")
    add_equation(doc, "Tⱼ = (wⱼᵀr) / sqrt(RSSᵣ/dfᵣ),     Tmax = max₁≤ⱼ≤q |Tⱼ|.", 5)
    add_para(doc, "For each resample, the entire maximum is recomputed. The resulting pmax is therefore familywise adjusted over the retained direction family. The unadjusted coordinate-level p-value, when returned, is labeled exploratory and must not be interpreted as confirmatory after selection.")

    doc.add_heading("4.5 Joint augmentation branch", level=2)
    add_para(doc, "The unrestricted design U = [B, W] yields residual sum of squares RSSu and residual degrees of freedom dfu. The partial-F statistic captures distributed misspecification whose energy is spread across several individually moderate coordinates.")
    add_equation(doc, "Fjoint = ((RSSᵣ − RSSᵤ)/q) / (RSSᵤ/dfᵤ).", 6)

    doc.add_heading("4.6 Residual resampling", level=2)
    add_para(doc, "The development implementation uses a frozen-direction reduced-model residual permutation. It forms ℓy*(b) = ŷ + P(b)r, where P(b) permutes residual entries, and evaluates the statistics with B and the selected W held fixed. With B-fixed structural directions, this is a standard reduced-model residual strategy whose accuracy depends on the exchangeability of residuals under the null. Because ŷ² and ŷ³ were constructed from the observed response and frozen across permutations, the fitted-value branch is response-adaptive; its finite-sample exactness is not guaranteed. The empirical p-values should therefore be read as calibrated approximations.")
    add_para(doc, "A confirmatory HX implementation should use one of two declared modes. Refit mode reconstructs the restricted fit, fitted-value directions, pivot selection, and both statistics inside every resample. Split mode constructs response-adaptive directions on training folds and tests them on disjoint observations, combining fold-level tests with a valid multiplicity correction. Under heteroskedasticity, unrestricted residual permutation should be replaced by a wild bootstrap; clustered or repeated observations require exchangeability blocks or cluster-level sign/weight schemes.")
    add_equation(doc, "p = (1 + Σᵦ 1{T*(b) ≥ Tobs}) / (R + 1),     b = 1,…,R.", 7)

    doc.add_heading("4.7 Combined decision and output contract", level=2)
    add_para(doc, "Let pmax be the adjusted maxT p-value and pF the resampling p-value for Fjoint. HX uses a Bonferroni minimum across these two branches:")
    add_equation(doc, "pHX = min{1, 2 min(pmax, pF)};     reject H₀ if pHX ≤ α.", 8)
    add_para(doc, "This combination does not require independence between branches. A result object contains pHX, pmax, pF, q, selected direction names, the maxT winner, observed statistics, sample size, ranks, degrees of freedom, conditioning diagnostics, resample count, seed, and a failure code. Returning only pHX is insufficient for scientific audit.")

    doc.add_heading("5 Statistical Interpretation and Safeguards", level=1)
    doc.add_heading("5.1 What error control is and is not claimed", level=2)
    add_para(doc, "The maxT branch controls search across the retained coordinate family when its resampling reference is valid. The final Bonferroni step controls the union of the maxT and joint-F branch decisions when both component p-values are valid. Neither statement repairs a misspecified resampling law. In the frozen fitted-value implementation, nominal size is evaluated empirically and the p-value is best described as approximate. The randomized null results are consistent with 5% calibration but are not a mathematical proof.")

    doc.add_heading("5.2 Invariance requirements", level=2)
    add_bullet(doc, "Translation invariance: adding a constant to ℓy changes only the intercept and must not trigger a different degeneracy classification.")
    add_bullet(doc, "Response-scale invariance: multiplying ℓy by a nonzero constant should leave standardized direction geometry and p-values unchanged up to numerical tolerance.")
    add_bullet(doc, "Predictor-scale invariance: affine recoding of log predictors is absorbed by standardization, except for transformations that change ordering or domain validity.")
    add_bullet(doc, "Column-order stability: a declared deterministic tie rule must make retained direction names reproducible when residual norms are numerically equal.")

    doc.add_heading("5.3 Guard conditions", level=2)
    add_table(doc, ["Condition", "Required response"], [
        ("X is not two-dimensional", "Return bad_shape_X; do not throw an opaque linear-algebra exception"),
        ("len(y) ≠ n", "Return length_mismatch"),
        ("Nonfinite input or log-domain violation", "Return nonfinite_input or domain_error"),
        ("rank(B) < columns(B)", "Return rank_deficient_base"),
        ("Condition number above threshold", "Return ill_conditioned_base with diagnostic value"),
        ("TSSℓy numerically zero", "Return constant_response"),
        ("RSSr/TSSℓy below tolerance", "Return degenerate_zero_residual"),
        ("No independent augmentation direction", "Return no_valid_direction"),
        ("dfu below floor", "Return insufficient_residual_df"),
    ], widths=[Inches(2.3), Inches(4.75)], font_size=8.8)

    doc.add_heading("6 Evaluation Design", level=1)
    doc.add_heading("6.1 Methods", level=2)
    add_para(doc, "Four methods were compared at α = 0.05: (i) a single-direction permutation form check; (ii) MD4, the structural-only direction library using maxT; (iii) HX, which adds fitted-value directions and the joint-F branch; and (iv) fitted-value RESET using its classical parametric F reference. The earlier manuscript also considered heteroskedasticity-consistent RESET variants, but the focused HX experiments used the classical fitted RESET comparator. All methods shared datasets, seeds, and decision thresholds; they did not share one common null generator. The permutation procedures used their declared empirical references, whereas RESET used its F reference. Power comparisons should therefore be interpreted jointly with each method’s empirical null calibration.")

    doc.add_heading("6.2 Controlled development benchmark", level=2)
    add_para(doc, "The controlled suite comprised three null scenarios and six structural alternatives—quadratic, cubic, interaction, dilution, opposing, and threshold departures—at n = 50 and n = 200, with selected cubic and opposing cases also evaluated at n = 1,000. Each principal cell used 50 seeded replications with 1% relative noise. Additional NIST-style analytic functions were used as interpretable challenge cases. This suite was used during method development and is not an untouched final test set.")

    doc.add_heading("6.3 Post-development randomized stress test", level=2)
    add_para(doc, "A master seed of 20261904 generated 1,000 groups: 393 null cases and 607 alternatives. Sample size, predictor count, noise, signal-to-noise ratio, exponents, and departure family varied across groups. The 12 alternative families were cross-cubic, cubic, exponential deformation, interaction, logistic bump, mixed cubic–oscillatory, opposing quadratic, quadratic, quartic, one-cycle sinusoid, two-cycle sinusoid, and threshold. The generator expands coverage, but its family design remained informed by development experience. We therefore call it a post-development randomized stress test rather than independent external validation.")

    doc.add_heading("6.4 Metrics and paired inference", level=2)
    add_para(doc, "For null cases, the primary metric was false-positive rate with a 95% Wilson interval. For alternatives, the primary metric was rejection rate (power) with the same interval. Overall accuracy used the generated class label only as a summary, not as a substitute for separate size and power. Paired differences on the same cases were evaluated with exact two-sided McNemar tests, equivalent to a binomial test on discordant pairs.")

    doc.add_heading("7 Results", level=1)
    doc.add_heading("7.1 Controlled benchmark", level=2)
    add_para(doc, "At n = 200, the pooled null rejection rates were 0.033 for the single-direction form check, 0.027 for MD4, 0.033 for HX, and 0.053 for RESET. In the six controlled alternative families, HX rejected all 50 replications at both n = 50 and n = 200. The main gains were complementary: RESET was comparatively weak for dilution and opposing alternatives, while the single-direction check was weak for cubic, interaction, and opposing alternatives.")
    add_caption(doc, "Table 1. Controlled alternative rejection rates by sample size")
    add_table(doc, ["Family", "n", "form_check", "MD4", "HX", "RESET"], [
        ("Quadratic", "50 / 200", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00"),
        ("Cubic", "50 / 200", "0.50 / 0.36", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00"),
        ("Interaction", "50 / 200", "0.24 / 0.18", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00"),
        ("Dilution", "50 / 200", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00", "0.44 / 0.88"),
        ("Opposing", "50 / 200", "0.00 / 0.02", "1.00 / 1.00", "1.00 / 1.00", "0.42 / 0.96"),
        ("Threshold", "50 / 200", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00", "1.00 / 1.00"),
    ], widths=[Inches(1.35), Inches(0.85), Inches(1.15), Inches(1.15), Inches(1.15), Inches(1.15)], font_size=8.5)

    doc.add_heading("7.2 Randomized calibration and power", level=2)
    doc.add_picture(str(perf_path), width=Inches(7.0))
    add_caption(doc, "Figure 2. False-positive rate and power in 1,000 randomized groups. The red line marks the nominal 0.05 level. Intervals are reported in Table 2.")
    metric_rows = []
    for name in ("form_check", "MD4", "HX", "RESET"):
        s = summary[name]
        lo0, hi0 = wilson(s["fp"], s["n0"])
        lo1, hi1 = wilson(s["tp"], s["n1"])
        acc = (s["n0"] - s["fp"] + s["tp"]) / (s["n0"] + s["n1"])
        metric_rows.append((name, f'{s["fp"]}/{s["n0"]}', f'{s["fpr"]:.4f} [{lo0:.4f}, {hi0:.4f}]', f'{s["tp"]}/{s["n1"]}', f'{s["tpr"]:.4f} [{lo1:.4f}, {hi1:.4f}]', f"{acc:.3f}"))
    add_caption(doc, "Table 2. Randomized stress test results with 95% Wilson intervals")
    add_table(doc, ["Method", "False positives", "FPR [95% CI]", "True positives", "Power [95% CI]", "Accuracy"], metric_rows,
              widths=[Inches(1.0), Inches(0.9), Inches(1.55), Inches(0.9), Inches(1.55), Inches(0.8)], font_size=8.2)
    add_para(doc, "HX achieved the highest power and overall accuracy while its estimated false-positive rate remained close to the nominal level. RESET’s false-positive estimate was 0.0611; its confidence interval still included 0.05. The relevant conclusion is therefore not that HX had a demonstrably lower size than every comparator, but that it delivered substantially greater power without evidence of size inflation in this generator.")

    doc.add_heading("7.3 Alternative-family coverage", level=2)
    family_rows = [
        ("cross_cubic", 52, .10, .29, 1.00, 1.00),
        ("cubic", 56, .41, 1.00, 1.00, .82),
        ("exp_deform", 44, 1.00, 1.00, 1.00, .91),
        ("interact", 47, .21, 1.00, 1.00, 1.00),
        ("logistic_bump", 49, .10, 1.00, 1.00, .78),
        ("mixed_cubic_osc", 54, .52, 1.00, 1.00, .87),
        ("opposing_quad", 58, .88, 1.00, 1.00, .05),
        ("quad", 58, 1.00, 1.00, 1.00, .98),
        ("quartic", 48, 1.00, 1.00, 1.00, 1.00),
        ("sin1", 53, .00, 1.00, 1.00, .66),
        ("sin2", 38, .18, .97, .97, .63),
        ("threshold", 50, .42, .44, .44, .38),
    ]
    add_caption(doc, "Table 3. Alternative rejection rate by randomized family")
    add_table(doc, ["Family", "n", "form_check", "MD4", "HX", "RESET"], family_rows,
              widths=[Inches(1.75), Inches(0.55), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1)], font_size=8.2)
    add_para(doc, "HX reached 1.00 rejection rate in ten of twelve families. HX and MD4 both detected 0.97 of the two-cycle sinusoid cases, while threshold departures remained difficult for every method. Cross-cubic alternatives illustrate the fitted-value contribution: MD4 detected 0.29, whereas HX and RESET both detected 1.00. Opposing quadratic alternatives illustrate the structural contribution: HX and MD4 detected all cases, but RESET detected only 0.05.")

    doc.add_heading("7.4 Paired discordance", level=2)
    add_caption(doc, "Table 4. Exact paired comparisons on the 607 randomized alternatives")
    add_table(doc, ["Comparator", "HX only", "Comparator only", "Exact two-sided p"], [
        ("form_check", 278, 0, "4.12 × 10⁻⁸⁴"),
        ("MD4", 37, 0, "1.46 × 10⁻¹¹"),
        ("RESET", 123, 1, "1.18 × 10⁻³⁵"),
    ], widths=[Inches(1.45), Inches(1.35), Inches(1.65), Inches(2.25)], font_size=9.0)
    add_para(doc, "On null cases, HX versus form_check had 11 versus 12 discordant rejections (p = 1.000), HX versus MD4 had 6 versus 2 (p = 0.289), and HX versus RESET had 9 versus 13 (p = 0.524). The paired evidence therefore supports a power difference under this alternative generator but not a meaningful size difference among the methods.")

    doc.add_heading("7.5 Direction attribution and analytic challenges", level=2)
    counts = Counter(r["argmax_hx"] for r in rows if r["klass"] == "H1" and r.get("argmax_hx"))
    total = sum(counts.values())
    top_rows = [(name, count, f"{100*count/total:.1f}%") for name, count in counts.most_common(8)]
    add_caption(doc, "Table 5. Most frequent winning HX directions among randomized alternatives")
    add_table(doc, ["Direction", "Count", "Share"], top_rows,
              widths=[Inches(3.2), Inches(1.2), Inches(1.2)], font_size=9.0)
    add_para(doc, "The fitted-value cubic direction was the most common winner, followed by structural cubic, quadratic, and oscillatory directions. This distribution supports the hybrid design: no single family accounted for all detections. It is descriptive, not a post-selection confidence statement about the true mechanism.")
    add_para(doc, "In the analytic challenge set, HX and RESET both alarmed on several nonlinear functions. A notable case was F3, where structural-only MD4 missed while fitted-value terms enabled HX to reject. Conversely, F10 produced a rank-deficiency failure for HX while RESET returned an alarm, demonstrating why failure states must remain visible. The L1 case had n = 16, only three unrestricted residual degrees of freedom, and a condition number near 4.9 × 10⁹; HX rejected while RESET did not, but the numerical regime makes the result inconclusive rather than a clean method victory.")

    doc.add_heading("7.6 Runtime", level=2)
    add_para(doc, "At n = 200 and p = 2, median development-run time was approximately 4.7 ms for HX, 4.6 ms for MD4, and 0.5 ms for parametric RESET on the recorded machine. These values are implementation- and hardware-dependent. Residual resampling dominates HX cost; confirmatory refit mode will be slower. Matrix-free projection B(B⁺D), batched resampling, and streaming maxima avoid unnecessary n × n projection matrices and full permutation-score storage.")

    doc.add_heading("8 Mechanistic Interpretation", level=1)
    doc.add_heading("8.1 Why the hybrid library helps", level=2)
    add_para(doc, "Structural directions supply locality and interpretability. A quadratic winner indicates curvature in one log predictor; an interaction winner indicates non-additivity; an oscillatory winner indicates a residual periodic pattern. These probes are weak when the omitted component lies outside the library or combines across predictors in a way that cancels a marginal projection. Fitted-value powers pool information along the restricted model’s predicted response and can capture broad departures without naming their source. They are correspondingly less diagnostic and response-adaptive.")
    add_para(doc, "The comparison between MD4 and HX is a structural ablation: adding fitted-value directions and the joint branch produced 37 HX-only versus no MD4-only rejection among randomized alternatives. It does not isolate the fitted-value library from the joint-F branch. A complete component ablation must cross structural-only versus fitted-only versus hybrid libraries with maxT-only versus joint-F-only versus combined decisions, using the same resampling policy and untouched confirmatory cases.")

    doc.add_heading("8.2 Concentrated versus distributed departures", level=2)
    add_para(doc, "maxT is suited to sparse departures: one direction carries most of the omitted signal. Joint F is suited to dense departures: several coordinates carry moderate signal. Combining them spends at most a factor of two in p-value but protects against committing to either geometry. This is a two-branch portfolio inside one test, analogous to the broader benchmark’s finding that structural and omnibus diagnostics are complementary.")

    doc.add_heading("8.3 Role in a model-discovery workflow", level=2)
    add_number(doc, "Generate candidate equations using physical constraints, symbolic regression, or expert hypotheses.")
    add_number(doc, "Reserve an untouched validation partition before tuning direction libraries or thresholds.")
    add_number(doc, "Fit each candidate on training data and apply HX under a declared confirmatory resampling mode on validation data.")
    add_number(doc, "Discard or revise candidates that HX falsifies; inspect winning directions to formulate the next hypothesis.")
    add_number(doc, "For candidates not rejected, report power diagnostics and the exact battery used; do not state that the equation has been proven.")

    doc.add_heading("9 Limitations and Threats to Validity", level=1)
    add_para(doc, "First, the evidence is not an independent external validation. The controlled suite informed method development, and the randomized family generator reflects the same problem-solving context. The strong power estimates may therefore be optimistic for unrelated scientific domains.")
    add_para(doc, "Second, the reported HX p-values use frozen fitted-value directions under reduced-model residual permutation. The resulting calibration is approximate because the directions are response-adaptive. Empirical null calibration was satisfactory in the available generator, but confirmatory use should reconstruct adaptive steps within resamples, use sample splitting, or justify an alternative resampling law.")
    add_para(doc, "Third, unrestricted residual permutation assumes exchangeability. Heteroskedastic, clustered, temporal, or spatial errors violate that condition. Wild bootstrap, exchangeability blocks, or domain-specific dependence models are required. Fourth, the fixed library has blind spots: the threshold family remained difficult, and no finite battery can cover all misspecification. Fifth, small residual degrees of freedom and ill-conditioned designs can make both directional selection and F statistics unstable. Sixth, the winning direction is post-selection and exploratory; causal or physical interpretation requires a separately specified follow-up analysis.")
    add_para(doc, "Finally, algorithmic novelty is narrower than scientific validity. HX combines established ideas—RESET-style fitted-value powers, structured basis augmentation, maxT adjustment, joint testing, and residual resampling—into an explicit falsification workflow. Its value must be judged by transparent error control, robustness, and out-of-domain performance rather than by naming the combination.")

    doc.add_heading("10 Reproducibility Protocol", level=1)
    add_para(doc, "A reproducible HX report should archive the raw positive variables, log-transformation rules, excluded observations, standardized design, restricted and augmented ranks, condition numbers, direction names, selection tolerance, residual-degree floor, α, resampling scheme, exchangeability blocks or bootstrap weights, resample count, random seed, observed Tmax and Fjoint, component and combined p-values, and failure code. The benchmark datasets and seeds must be immutable once confirmatory evaluation begins.")
    add_para(doc, "For a submission-quality replication, use at least 9,999 resamples near the decision boundary, rerun with independent seeds, verify the four invariance properties in Section 5.2, and include regression tests for malformed shapes, length mismatch, constant response, exact fit, rank deficiency, and ill conditioning. Compare matrix-free and reference implementations on small problems to numerical tolerance. Report every attempted case, including failures, rather than deleting cases without a p-value.")
    add_table(doc, ["Evidence tier", "Permitted claim"], [
        ("Development benchmark", "The method was tuned and its mechanisms were explored on these cases"),
        ("Post-development randomized stress test", "The method generalized across randomized variants of declared families"),
        ("Locked synthetic confirmation", "Size and power were estimated on unseen seeds and preregistered families"),
        ("External scientific datasets", "Transportability was assessed in domains not used to design the library"),
    ], widths=[Inches(2.1), Inches(4.95)], font_size=9.0)

    doc.add_heading("11 Conclusion", level=1)
    add_para(doc, "HX is a hybrid specification test for falsifying proposed log-linear physical models. It combines interpretable structural directions with fitted-value powers, removes redundant directions, evaluates concentrated and distributed departures, and records numerical failure states. In the available randomized stress test, HX paired near-nominal null rejection with substantially higher power than a single-direction form check, structural-only MD4, and classical RESET. The result supports further evaluation, not unconditional deployment. The central methodological lesson is that a credible model check must expose its tested directions, multiplicity correction, resampling assumptions, and failures. Used in that form, HX can serve as a rigorous challenge layer in scientific model discovery.")

    doc.add_page_break()
    doc.add_heading("Appendix A Algorithm Specification", level=1)
    add_para(doc, "Algorithm 1. HX with a declared resampling mode")
    algorithm_steps = [
        "Validate that Xℓ is finite and two-dimensional, ℓy is finite and aligned, n exceeds the residual-degree floor, and requested parameters are valid.",
        "Standardize columns of Xℓ, form B = [1, Z], check rank and conditioning, and fit the restricted model.",
        "Compute ŷ and r; reject constant-response or degenerate-exact-fit cases using TSS-based criteria.",
        "Construct named structural directions and standardized fitted-value powers.",
        "Residualize candidates against B without materializing an n × n projector.",
        "Use modified Gram–Schmidt with deterministic, scale-aware pivoting to retain q independent coordinates W.",
        "Compute observed Tmax and Fjoint; record the winning direction as exploratory.",
        "Generate R resamples under the declared mode. For confirmatory response-adaptive inference, rebuild every adaptive step inside each resample or use a split-sample construction.",
        "Compute pmax and pF with the plus-one correction, then pHX = min(1, 2 min(pmax, pF)).",
        "Return the complete result object or a named failure state; never return a silent NaN or an unclassified exception.",
    ]
    for idx, step in enumerate(algorithm_steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(f"{idx}.  ").bold = True
        p.add_run(step)

    doc.add_heading("Appendix B Failure Taxonomy", level=1)
    add_table(doc, ["Class", "Examples", "Interpretation"], [
        ("Input", "bad_shape_X, length_mismatch, nonfinite_input", "The requested test is undefined for the supplied arrays"),
        ("Domain", "domain_error", "Log transformation or physical-domain condition failed"),
        ("Base model", "rank_deficient_base, ill_conditioned_base", "The restricted representation is not numerically identifiable"),
        ("Response", "constant_response, degenerate_zero_residual", "There is insufficient residual variation for a specification test"),
        ("Augmentation", "no_valid_direction, insufficient_residual_df", "The added library contributes no stable testable dimension"),
        ("Resampling", "invalid_exchangeability, resampling_failure", "The declared null generator is unjustified or could not be executed"),
    ], widths=[Inches(1.1), Inches(2.4), Inches(3.55)], font_size=8.8)

    doc.add_heading("Appendix C Minimum Reporting Checklist", level=1)
    checklist = [
        "State the restricted equation, transformations, and positivity/domain rules.",
        "List every candidate direction and whether it is response-adaptive.",
        "Report retained directions, ranks, conditioning, residual degrees of freedom, and tolerances.",
        "Name the resampling or bootstrap scheme and justify exchangeability or symmetry assumptions.",
        "Report R, seeds, component p-values, combined p-value, α, and all failure states.",
        "Label the winning direction and any unadjusted coordinate p-value as exploratory.",
        "Separate development, stress-test, confirmatory, and external evidence.",
        "Archive generator code and immutable confirmatory seeds before final evaluation.",
    ]
    for item in checklist:
        add_bullet(doc, item)

    doc.add_heading("References", level=1)
    refs = [
        "[1] Ramsey, J. B. (1969). Tests for specification errors in classical linear least-squares regression analysis. Journal of the Royal Statistical Society Series B, 31, 350–371. https://doi.org/10.1111/j.2517-6161.1969.tb00796.x",
        "[2] MacKinnon, J. G., & White, H. (1985). Some heteroskedasticity-consistent covariance matrix estimators with improved finite sample properties. Journal of Econometrics, 29, 305–325. https://doi.org/10.1016/0304-4076(85)90158-7",
        "[3] Freedman, D., & Lane, D. (1983). A nonstochastic interpretation of reported significance levels. Journal of Business & Economic Statistics, 1, 292–298. https://doi.org/10.1080/07350015.1983.10509354",
        "[4] Winkler, A. M., Ridgway, G. R., Webster, M. A., Smith, S. M., & Nichols, T. E. (2014). Permutation inference for the general linear model. NeuroImage, 92, 381–397. https://doi.org/10.1016/j.neuroimage.2014.01.060",
        "[5] Davidson, R., & Flachaire, E. (2008). The wild bootstrap, tamed at last. Journal of Econometrics, 146, 162–169. https://doi.org/10.1016/j.jeconom.2008.08.003",
        "[6] Westfall, P. H., & Young, S. S. (1993). Resampling-Based Multiple Testing: Examples and Methods for P-Value Adjustment. Wiley.",
        "[7] Bakarji, J., Callaham, J. L., Brunton, S. L., & Kutz, J. N. (2022). Dimensionally consistent learning with Buckingham Pi. Nature Computational Science, 2, 834–844. https://doi.org/10.1038/s43588-022-00355-5",
        "[8] La Cava, W., Orzechowski, P., Burlacu, B., et al. (2021). Contemporary symbolic regression methods and their relative performance. Proceedings of NeurIPS Datasets and Benchmarks.",
        "[9] Udrescu, S.-M., & Tegmark, M. (2020). AI Feynman: A physics-inspired method for symbolic regression. Science Advances, 6, eaay2631. https://doi.org/10.1126/sciadv.aay2631",
        "[10] Cranmer, M. (2023). Interpretable machine learning for science with PySR and SymbolicRegression.jl. arXiv:2305.01582. https://doi.org/10.48550/arXiv.2305.01582",
        "[11] Matsubara, Y., Chiba, N., Igarashi, R., & Ushiku, Y. (2022). Rethinking symbolic regression datasets and benchmarks for scientific discovery. arXiv:2206.10540. https://doi.org/10.48550/arXiv.2206.10540",
        "[12] Fisher, R. A. (1935). The Design of Experiments. Oliver & Boyd.",
        "[13] White, H. (1980). A heteroskedasticity-consistent covariance matrix estimator and a direct test for heteroskedasticity. Econometrica, 48, 817–838. https://doi.org/10.2307/1912934",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ref)
        r.font.size = Pt(9)

    # Ensure all sections share page settings and headers after additions.
    for section in doc.sections:
        setup_page(section)
        if not section.header.paragraphs[0].text:
            add_header_footer(section)

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_document()
